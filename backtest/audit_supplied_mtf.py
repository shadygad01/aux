"""Streaming and lineage audit for owner-supplied XAUUSD H1/M15/M5 files."""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from collections.abc import Iterator, Sequence
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

PRICE_SCALE = 100_000.0


@dataclass(frozen=True, slots=True)
class Bar:
    timestamp: datetime
    open: float
    high: float
    low: float
    close: float
    volume_bid: float
    volume_ask: float


def _xlsx_rows(path: Path) -> Iterator[tuple[Any, ...]]:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    yield from sheet.iter_rows(values_only=True)
    workbook.close()


def _csv_rows(path: Path) -> Iterator[tuple[Any, ...]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        yield from (tuple(row) for row in csv.reader(handle))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load(path: Path) -> tuple[tuple[str, ...], list[Bar], set[str]]:
    rows = _xlsx_rows(path) if path.suffix.lower() == ".xlsx" else _csv_rows(path)
    header = tuple(str(value) for value in next(rows))
    bars: list[Bar] = []
    sources: set[str] = set()
    for row in rows:
        if not any(value is not None and value != "" for value in row):
            continue
        timestamp = datetime.fromisoformat(str(row[0]).replace("Z", "+00:00")).astimezone(UTC)
        bars.append(
            Bar(
                timestamp=timestamp,
                open=float(row[1]) / PRICE_SCALE,
                high=float(row[2]) / PRICE_SCALE,
                low=float(row[3]) / PRICE_SCALE,
                close=float(row[4]) / PRICE_SCALE,
                volume_bid=float(row[5]),
                volume_ask=float(row[6]),
            )
        )
        sources.add(str(row[7]))
    return header, bars, sources


def audit(
    path: Path, bars: Sequence[Bar], header: Sequence[str], sources: set[str]
) -> dict[str, object]:
    timestamps = [bar.timestamp for bar in bars]
    deltas = Counter(
        int((right - left).total_seconds())
        for left, right in zip(timestamps, timestamps[1:], strict=False)
    )
    duplicates = len(timestamps) - len(set(timestamps))
    invalid = sum(
        bar.high < max(bar.open, bar.close)
        or bar.low > min(bar.open, bar.close)
        or bar.high < bar.low
        for bar in bars
    )
    zero_volume = sum(bar.volume_bid == 0 and bar.volume_ask == 0 for bar in bars)
    flat_zero = sum(
        bar.volume_bid == 0 and bar.volume_ask == 0 and bar.open == bar.high == bar.low == bar.close
        for bar in bars
    )
    return {
        "path": str(path),
        "sha256": _sha256(path),
        "bytes": path.stat().st_size,
        "header": list(header),
        "sources": sorted(sources),
        "rows": len(bars),
        "first": timestamps[0].isoformat(),
        "last": timestamps[-1].isoformat(),
        "duplicates": duplicates,
        "strictly_increasing": all(
            right > left for left, right in zip(timestamps, timestamps[1:], strict=False)
        ),
        "top_deltas_seconds": deltas.most_common(8),
        "invalid_ohlc": invalid,
        "zero_volume_bars": zero_volume,
        "flat_zero_volume_bars": flat_zero,
        "minimum_price": min(bar.low for bar in bars),
        "maximum_price": max(bar.high for bar in bars),
    }


def _aggregate(bars: Sequence[Bar], seconds: int) -> dict[datetime, Bar]:
    groups: dict[datetime, list[Bar]] = {}
    for bar in bars:
        epoch = int(bar.timestamp.timestamp())
        timestamp = datetime.fromtimestamp(epoch - epoch % seconds, UTC)
        groups.setdefault(timestamp, []).append(bar)
    return {
        timestamp: Bar(
            timestamp,
            group[0].open,
            max(bar.high for bar in group),
            min(bar.low for bar in group),
            group[-1].close,
            sum(bar.volume_bid for bar in group),
            sum(bar.volume_ask for bar in group),
        )
        for timestamp, group in groups.items()
    }


def lineage(base: Sequence[Bar], supplied: Sequence[Bar], seconds: int) -> dict[str, object]:
    derived = _aggregate(base, seconds)
    supplied_by_time = {bar.timestamp: bar for bar in supplied}
    overlap = sorted(derived.keys() & supplied_by_time.keys())
    price_mismatch = volume_mismatch = 0
    examples: list[dict[str, object]] = []
    for timestamp in overlap:
        left, right = derived[timestamp], supplied_by_time[timestamp]
        price_equal = all(
            abs(a - b) <= 1e-8
            for a, b in zip(
                (left.open, left.high, left.low, left.close),
                (right.open, right.high, right.low, right.close),
                strict=False,
            )
        )
        volume_equal = (
            abs(left.volume_bid - right.volume_bid) <= 1e-6
            and abs(left.volume_ask - right.volume_ask) <= 1e-6
        )
        price_mismatch += not price_equal
        volume_mismatch += not volume_equal
        if (not price_equal or not volume_equal) and len(examples) < 5:
            examples.append(
                {
                    "timestamp": timestamp.isoformat(),
                    "derived": asdict(left),
                    "supplied": asdict(right),
                }
            )
    return {
        "seconds": seconds,
        "derived_groups": len(derived),
        "supplied_bars": len(supplied),
        "overlap": len(overlap),
        "missing_from_supplied": len(derived.keys() - supplied_by_time.keys()),
        "extra_in_supplied": len(supplied_by_time.keys() - derived.keys()),
        "price_mismatches": price_mismatch,
        "volume_mismatches": volume_mismatch,
        "examples": examples,
    }


def main(arguments: Sequence[str]) -> dict[str, object]:
    m5_path, m15_path, h1_path = map(Path, arguments)
    m5_header, m5, m5_sources = load(m5_path)
    m15_header, m15, m15_sources = load(m15_path)
    h1_header, h1, h1_sources = load(h1_path)
    return {
        "schema_version": "1.0.0",
        "price_scale_applied": PRICE_SCALE,
        "files": {
            "M5": audit(m5_path, m5, m5_header, m5_sources),
            "M15": audit(m15_path, m15, m15_header, m15_sources),
            "H1": audit(h1_path, h1, h1_header, h1_sources),
        },
        "lineage": {
            "M5_to_M15": lineage(m5, m15, 900),
            "M5_to_H1": lineage(m5, h1, 3600),
        },
        "audited_at": datetime.now(UTC).isoformat(),
    }


if __name__ == "__main__":
    import sys

    result = json.dumps(main(sys.argv[1:4]), indent=2, default=str)
    if len(sys.argv) > 4:
        output = Path(sys.argv[4])
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(result, encoding="utf-8")
    print(result)
